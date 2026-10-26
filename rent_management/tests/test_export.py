# Copyright (c) 2026, Shashank Dixit and Contributors
# See license.txt

from io import BytesIO

import frappe
from frappe.tests import IntegrationTestCase
from openpyxl import load_workbook

from rent_management import export
from rent_management.rent_management_system.doctype.tenant.test_tenant import make_aadhaar
from rent_management.rent_management_system.doctype.tenant_space_allotment.test_tenant_space_allotment import (
	make_allotment,
	make_room,
)


class TestExport(IntegrationTestCase):
	def test_workbook(self):
		aadhaar = make_aadhaar("56789012345")
		tenant = frappe.get_doc(
			{"doctype": "Tenant", "first_name": "Export", "last_name": "Me", "aadhar_card_number": aadhaar}
		).insert()
		make_allotment(room=make_room(), tenant=tenant)

		workbook = load_workbook(BytesIO(export.get_workbook()), read_only=True)
		self.assertEqual(workbook.sheetnames, [title for title, *_rest in export.SHEETS])
		tenants = list(workbook["Tenants"].values)
		self.assertIn("First Name", tenants[0])
		row = next(r for r in tenants if r[0] == tenant.name or str(r[0]) == str(tenant.name))
		self.assertIn(aadhaar[-4:], row)
		# the full Aadhaar number never leaves the server
		for sheet in workbook.worksheets:
			for values in sheet.values:
				self.assertNotIn(aadhaar, [str(v) for v in values])
